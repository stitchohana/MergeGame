"""Round-trip the dedicated golden-finger workbook without changing other tables."""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = {
    "artifacts": ["id", "name", "unlock_level", "unlock_label", "daily_free", "daily_buy_limit", "base_cost", "cost_increment", "description"],
    "realms": ["cultivation_min", "cultivation_max", "level_cap"],
    "settings": ["reset_hour", "rewind_seconds"],
}
TEXT = {"id", "name", "unlock_label", "description"}


def validate(data):
    if [row["id"] for row in data["artifacts"]] != ["basin", "furnace", "sand"]:
        raise ValueError("golden_finger: expected basin, furnace, sand")
    for sheet, columns in COLUMNS.items():
        if not data[sheet]:
            raise ValueError(f"golden_finger.{sheet}: empty")
        for row in data[sheet]:
            for key in columns:
                value = row[key]
                if key in TEXT:
                    if not isinstance(value, str) or not value.strip():
                        raise ValueError(f"golden_finger.{key}: expected text")
                elif isinstance(value, bool) or not isinstance(value, int) or value < (0 if key == "reset_hour" else 1):
                    raise ValueError(f"golden_finger.{key}: expected integer")
    if len(data["settings"]) != 1 or data["settings"][0]["reset_hour"] > 23:
        raise ValueError("golden_finger: invalid settings")
    previous = 0
    for row in data["realms"]:
        if row["cultivation_min"] != previous + 1 or row["cultivation_max"] < row["cultivation_min"]:
            raise ValueError("golden_finger: invalid realm ranges")
        previous = row["cultivation_max"]
    return data


def read_golden_finger(workbook):
    data = {}
    for name, columns in COLUMNS.items():
        rows = list(workbook[name].values)
        if list(rows[0]) != columns:
            raise ValueError(f"golden_finger.{name}: invalid columns")
        data[name] = []
        for values in rows[1:]:
            if all(value is None for value in values):
                continue
            row = dict(zip(columns, values))
            for key in columns:
                if key not in TEXT:
                    value = row[key]
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or value != int(value):
                        raise ValueError(f"golden_finger.{key}: invalid number")
                    row[key] = int(value)
            data[name].append(row)
    return validate(data)


def validate_cultivation(data, stages):
    validate(data)
    if data["realms"][-1]["cultivation_max"] < len(stages):
        raise ValueError("golden_finger: realm ranges do not cover cultivation stages")
    for artifact in data["artifacts"]:
        level = artifact["unlock_level"]
        if level > len(stages) or stages[level - 1]["name"] != artifact["unlock_label"]:
            raise ValueError(f"golden_finger.{artifact['id']}: unlock level/name mismatch")
    return data


def write_golden_finger(data, path):
    validate(data)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, columns in COLUMNS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in data[name]:
            sheet.append([row[key] for key in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="326459")
        for index, key in enumerate(columns, 1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(index)].width = 68 if key == "description" else 22
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 64 if name == "artifacts" else 28
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    workbook.save(path)


if __name__ == "__main__":
    base = Path(__file__).parent
    write_golden_finger(json.loads((base / "json_output/golden_finger.json").read_text(encoding="utf-8")), base / "xlsx/golden_finger.xlsx")
    print("GOLDEN_FINGER_WORKBOOK_OK")
