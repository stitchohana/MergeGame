"""Validate the shipped realm workbook and the actual Excel/JSON schema adapter."""
import copy
import json
import sys
from pathlib import Path
from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "config"))
from secret_realm_table import SECRET_REALM_COLUMNS, read_secret_realm, secret_realm_rows, validate_secret_realm

data = json.loads((ROOT / "config/json_output/secret_realm.json").read_text(encoding="utf-8"))
assert read_secret_realm(load_workbook(ROOT / "config/xlsx/secret_realm.xlsx")) == data
assert data["order_supply"] == [{"id": "sword_1", "count": 3, "immovable": False}]
assert data["order_supply_realms"] == [
    {"cultivation_min": low, "cultivation_max": high,
     "drops": [{"id": "sword_" + str(n), "count": 3, "immovable": False} for n in range(1, cap + 1)]}
    for low, high, cap in [(11, 13, 2), (14, 16, 3), (17, 19, 4)]
]
fixed = [c for c in data["layout"] if c["id"].startswith("sword_") and c["immovable"]]
assert len(data["stages"]) == 6 and len(data["layout"]) == 6 * 28
assert {c["id"] for c in fixed} == {"sword_" + str(n) for n in range(1, 11)}
assert all(c["id"] == "sword_" + str(min(11, c["row"] // 2 + 1) - (1 if c["row"] >= 12 else 0)) for c in fixed)
for index, stage in enumerate(data["stages"], 1):
    boss = next(m for m in data["monsters"] if m["id"] == stage["boss"])
    assert boss["costs"] == [{"id": "sword_" + str(index + 6), "count": index}, {"id": "qi_6", "count": index if index >= 3 else 1}]
    mob = next(m for m in data["monsters"] if m["id"] == "mob_" + str(index))
    assert mob["costs"] == [{"id": "sword_" + str(index + 3 + (1 if index >= 3 else 0)), "count": 1}, {"id": "qi_" + str(min(6, max(3, index + 1))), "count": 1}]
qi = [c for c in data["layout"] if c["id"].startswith("qi_")]
assert len(qi) == 17 and all(c["fog"] and c["immovable"] for c in qi)
assert all(sum(c["id"] == "qi_" + str(n) for c in qi) == (9 if n == 5 else 2) for n in range(1, 6))
assert next(c for c in data["layout"] if (c["col"], c["row"]) == (2, 5))["id"] == "qi_5"
assert [(r["min_level"], r["max_level"], d["id"], d["count"], d["immovable"]) for r in data["merge_drops"] for d in r["drops"] if d["id"].startswith("qi_")] == [
    (5, 6, "qi_1", 1, False), (7, 8, "qi_2", 1, False),
    (9, 10, "qi_3", 1, False), (11, 12, "qi_4", 1, False),
]
assert [(r["min_level"], r["max_level"], d["id"], d["count"], d["immovable"]) for r in data["merge_drops"] for d in r["drops"] if d["id"].startswith("sword_")] == [
    (5 + tier * 2, 6 + tier * 2, "sword_" + str(level), 1, True)
    for tier in range(4) for level in range(max(1, tier), tier + 2)
]
# Exercise the same row exporter used by json_to_xlsx without touching other tables.
book = Workbook()
book.remove(book.active)
for name, columns in SECRET_REALM_COLUMNS.items():
    sheet = book.create_sheet(name)
    sheet.append(columns)
    for row in secret_realm_rows(data, name):
        sheet.append(row)
assert read_secret_realm(book) == data
activities = list(load_workbook(ROOT / "config/xlsx/activities.xlsx")["activities"].values)
realm_rows = [dict(zip(activities[0], row)) for row in activities[1:] if row[0] == 9]
assert len(realm_rows) == 1 and realm_rows[0]["widget"] == "SecretRealm"


def rejects(edit):
    invalid = copy.deepcopy(data)
    edit(invalid)
    try:
        validate_secret_realm(invalid)
    except ValueError:
        return
    raise AssertionError("invalid secret-realm config accepted")


rejects(lambda d: d["items"][0].update(level=99))
rejects(lambda d: d["order_supply_realms"][0].update(cultivation_min=0))
rejects(lambda d: d["order_supply_realms"][1].update(cultivation_min=13))
rejects(lambda d: d["order_supply_realms"][0]["drops"][0].update(immovable=True))
rejects(lambda d: d["order_supply_realms"][0]["drops"][0].update(id="qi_1"))
rejects(lambda d: d["order_supply_realms"][0].update(drops=[]))
rejects(lambda d: d["layout"][1].update(col=0))
rejects(lambda d: d["layout"][0].update(immovable=True))
rejects(lambda d: d["layout"][0].update(fog=True))
rejects(lambda d: d["layout"][3].update(id="stamina_1"))
rejects(lambda d: d["monsters"][1]["costs"][1].update(id="qi_5"))
rejects(lambda d: next(drop for rule in d["merge_drops"] for drop in rule["drops"] if drop["id"] == "sword_1").update(immovable=False))
rejects(lambda d: next(drop for rule in d["merge_drops"] for drop in rule["drops"] if drop["id"] == "sword_1").update(count=3))
rejects(lambda d: d["merge_drops"].append({"min_level": 7, "max_level": 8, "drops": [{"id": "sword_1", "count": 1, "immovable": True}]}))
rejects(lambda d: d["monsters"][0]["drops"][0].update(id="sword_1"))
rejects(lambda d: d["stages"][0].update(boss="boss_3"))
rejects(lambda d: next(c for c in d["layout"] if (c["col"], c["row"]) == (2, 5)).update(id="", monster="mob_1"))
rejects(lambda d: next(c for c in d["layout"] if (c["col"], c["row"]) == (2, 5)).update(fog=False))
rejects(lambda d: [c.update(fog=True) for c in d["layout"]])
rejects(lambda d: next(c for c in d["layout"] if c["row"] == 2 and c["id"]).update(id="sword_8"))
print("SECRET_REALM_CONFIG_SMOKE_OK")
